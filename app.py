import base64
import html as htmllib
import io
import json
import re
from collections import Counter
from pathlib import Path

import streamlit as st

st.set_page_config(page_title="Text Summarizer · Study Sparks AI", page_icon="✦", layout="wide")

OWNER_NAME = "Ubaid ur Rehman"


# ---------- Logo ----------
def logo_tag():
    mimes = {".png": "image/png", ".jpg": "image/jpeg", ".jpeg": "image/jpeg",
             ".webp": "image/webp", ".svg": "image/svg+xml"}
    for ext, mime in mimes.items():
        p = Path(__file__).parent / f"logo{ext}"
        if p.exists():
            b64 = base64.b64encode(p.read_bytes()).decode()
            return f'<img class="logo" src="data:{mime};base64,{b64}" alt="Study Sparks AI logo">'
    return '<span class="spark">✦</span>'  # fallback if no logo file found


# ---------- Styling ----------
st.markdown("""
<style>
@import url('https://fonts.googleapis.com/css2?family=Bricolage+Grotesque:wght@600;800&family=DM+Sans:wght@400;500;700&display=swap');
html, body, [class*="css"] { font-family: 'DM Sans', sans-serif; }
#MainMenu, footer, header { visibility: hidden; }
.block-container { padding-top: 1.2rem; max-width: 1150px; }
.topbar { display:flex; justify-content:space-between; align-items:center; flex-wrap:wrap; gap:12px;
          padding-bottom:14px; border-bottom:1px solid #E3DFF7; }
.brand { display:flex; align-items:center; gap:12px; font-family:'Bricolage Grotesque',sans-serif;
         font-weight:800; font-size:1.45rem; color:#1B1740; }
.brand .logo { height:46px; width:46px; object-fit:contain; border-radius:10px; }
.brand .spark { color:#FFB627; font-size:1.6rem; }
.brand b { color:#6B4EFF; }
.owner { background:#1B1740; color:#F6F4FF; padding:8px 16px; border-radius:999px; font-weight:500; font-size:.95rem; }
.owner b { color:#FFB627; }
.hero h1 { font-family:'Bricolage Grotesque',sans-serif; font-weight:800; font-size:2.5rem;
           line-height:1.1; color:#1B1740; margin:1.2rem 0 .3rem; letter-spacing:-.02em; }
.hero p { color:#4B466F; font-size:1.1rem; margin-bottom:1.2rem; }
.summary-box { background:#fff; border-left:6px solid #FFB627; border-radius:14px; padding:18px 20px;
               font-size:1.05rem; line-height:1.75; color:#1B1740; box-shadow:0 14px 30px -22px rgba(27,23,64,.5); }
.chip { display:inline-block; background:rgba(107,78,255,.12); color:#6B4EFF; padding:4px 12px;
        border-radius:999px; font-weight:700; font-size:.88rem; margin:0 6px 6px 0; }
div.stButton > button:first-child, div.stDownloadButton > button {
    background:#6B4EFF; color:#fff; border:0; border-radius:12px; padding:.65rem 1.2rem; font-weight:700; }
div.stButton > button:first-child:hover, div.stDownloadButton > button:hover { background:#5237E0; color:#fff; }
[data-testid="stMetricValue"] { color:#23B89A; font-family:'Bricolage Grotesque',sans-serif; }
textarea { border-radius:14px !important; }
</style>
""", unsafe_allow_html=True)

st.markdown(f"""
<div class="topbar">
  <div class="brand">{logo_tag()}<span>Study Sparks <b>AI</b></span></div>
  <div class="owner">Created by <b>{OWNER_NAME}</b></div>
</div>
<div class="hero">
  <h1>Turn long notes into a summary you can revise from.</h1>
  <p>Paste text or upload PDF, Word, PowerPoint, Excel, images and more. Pick a length. Get the key points.</p>
</div>
""", unsafe_allow_html=True)

# ---------- File reading ----------
FILE_TYPES = ["pdf", "docx", "pptx", "xlsx", "xlsm", "csv", "txt", "md", "html", "htm",
              "rtf", "json", "png", "jpg", "jpeg", "webp", "bmp", "tiff"]


@st.cache_data(show_spinner=False)
def extract_text(name: str, data: bytes) -> str:
    ext = Path(name).suffix.lower()

    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(io.BytesIO(data))
        return "\n".join((page.extract_text() or "") for page in reader.pages)

    if ext == ".docx":
        import docx
        d = docx.Document(io.BytesIO(data))
        parts = [p.text for p in d.paragraphs]
        for table in d.tables:
            for row in table.rows:
                parts.append(" ".join(c.text for c in row.cells))
        return "\n".join(parts)

    if ext == ".pptx":
        from pptx import Presentation
        prs = Presentation(io.BytesIO(data))
        parts = []
        for slide in prs.slides:
            for shape in slide.shapes:
                if getattr(shape, "has_text_frame", False) and shape.has_text_frame:
                    parts.append(shape.text_frame.text)
            if slide.has_notes_slide:
                parts.append(slide.notes_slide.notes_text_frame.text)
        return "\n".join(parts)

    if ext in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook
        wb = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        parts = []
        for ws in wb.worksheets:
            parts.append(f"{ws.title}.")
            for row in ws.iter_rows(values_only=True):
                cells = [str(c) for c in row if c is not None]
                if cells:
                    parts.append(", ".join(cells) + ".")
        return "\n".join(parts)

    if ext in (".png", ".jpg", ".jpeg", ".webp", ".bmp", ".tiff"):
        from PIL import Image
        import pytesseract
        return pytesseract.image_to_string(Image.open(io.BytesIO(data)))

    text = data.decode("utf-8", errors="ignore")

    if ext in (".html", ".htm"):
        text = re.sub(r"(?is)<(script|style).*?>.*?</\1>", " ", text)
        text = re.sub(r"(?s)<[^>]+>", " ", text)
        return htmllib.unescape(text)

    if ext == ".rtf":
        text = re.sub(r"\\'[0-9a-f]{2}", " ", text)
        text = re.sub(r"\\[a-z]+-?\d* ?", " ", text)
        return re.sub(r"[{}]", " ", text)

    if ext == ".json":
        try:
            vals = []

            def walk(o):
                if isinstance(o, dict):
                    for v in o.values():
                        walk(v)
                elif isinstance(o, list):
                    for v in o:
                        walk(v)
                elif isinstance(o, str):
                    vals.append(o)
            walk(json.loads(text))
            return "\n".join(vals)
        except Exception:
            return text

    return text  # txt, md, csv and anything else readable


# ---------- Summarizer logic ----------
STOP = set("""a about above after again all also am an and any are as at be because been before being below
between both but by can could did do does doing down during each few for from further had has have having he her
here hers him his how i if in into is it its just me more most my no nor not now of off on once only or other our
out over own same she should so some such than that the their them then there these they this those through to too
under until up very was we were what when where which while who whom why will with would you your""".split())

LEVELS = {"Short": (0.15, 2, 4), "Medium": (0.30, 3, 8), "Detailed": (0.45, 4, 14)}
MIN_WORDS = 40


def words_of(text):
    return re.findall(r"[\w'’-]+", text)


def split_sentences(text):
    text = re.sub(r"\s+", " ", text).strip()
    parts = re.split(r"(?<=[.!?])\s+", text)
    return [p for p in parts if len(words_of(p)) >= 4]


def local_summarize(text, level):
    sentences = split_sentences(text)
    freq = Counter(w for w in (x.lower() for x in words_of(text))
                   if w not in STOP and len(w) > 2 and not w.isdigit())
    if not freq or not sentences:
        return []
    top = max(freq.values())
    scored = []
    for i, s in enumerate(sentences):
        ws = [w.lower() for w in words_of(s)]
        score = sum(freq[w] / top for w in ws if w in freq) / (len(ws) ** 0.5)
        if i == 0:
            score *= 1.35
        elif i == len(sentences) - 1:
            score *= 1.1
        if len(ws) > 45:
            score *= 0.75
        if re.search(r"\b(important|key|main|therefore|in conclusion|significant|defined as|means)\b", s, re.I):
            score *= 1.15
        scored.append((score, i, s))
    ratio, lo, hi = LEVELS[level]
    n = min(hi, max(lo, round(len(sentences) * ratio)))
    best = sorted(scored, reverse=True)[:n]
    return [s for _, _, s in sorted(best, key=lambda x: x[1])]


def top_keywords(text, n=8):
    f = Counter(w for w in (x.lower() for x in words_of(text))
                if w not in STOP and len(w) > 3 and not w.isdigit())
    return [w for w, _ in f.most_common(n)]


def ai_summarize(text, level, fmt):
    import anthropic
    client = anthropic.Anthropic(api_key=st.secrets["ANTHROPIC_API_KEY"])
    size = {"Short": "2-3 sentences", "Medium": "4-6 sentences", "Detailed": "8-10 sentences"}[level]
    style = ('Return only bullet points, one per line, each starting with "- ".'
             if fmt == "Bullet points" else "Return a single clear paragraph with no heading.")
    prompt = (f"Summarize the following study material for a student. Keep it accurate and simple, in the "
              f"original language. Length: about {size}. {style} No commentary.\n\n<text>\n{text[:60000]}\n</text>")
    msg = client.messages.create(model="claude-sonnet-5-5", max_tokens=1000,
                                 messages=[{"role": "user", "content": prompt}])
    out = msg.content[0].text.strip()
    return [re.sub(r"^[-•*]\s*", "", l).strip() for l in out.splitlines() if l.strip()]


def ai_available():
    try:
        return "ANTHROPIC_API_KEY" in st.secrets
    except Exception:
        return False


SAMPLE = ("Photosynthesis is the process by which green plants, algae and some bacteria convert light energy into "
          "chemical energy. It takes place mainly in the chloroplasts, which contain a green pigment called "
          "chlorophyll. Chlorophyll absorbs sunlight, especially in the red and blue wavelengths, and reflects "
          "green light, which is why leaves look green. The process has two main stages. In the light-dependent "
          "reactions, sunlight splits water molecules, releasing oxygen and producing the energy carriers ATP and "
          "NADPH. In the Calvin cycle, the plant uses ATP and NADPH to turn carbon dioxide from the air into "
          "glucose. Glucose is then used for growth, stored as starch, or broken down during cellular respiration "
          "to release energy. Light intensity, carbon dioxide concentration and temperature all affect the rate of "
          "photosynthesis. If any of these is in short supply it becomes a limiting factor and slows the whole "
          "process. Photosynthesis is important for life on Earth because it produces the oxygen we breathe, forms "
          "the base of nearly every food chain, and removes carbon dioxide from the atmosphere.")

st.session_state.setdefault("text", "")
st.session_state.setdefault("result", None)
st.session_state.setdefault("up_key", 0)


def load_sample():
    st.session_state.text = SAMPLE


def clear_all():
    st.session_state.text = ""
    st.session_state.result = None
    st.session_state.up_key += 1  # resets the file uploader


# ---------- Layout ----------
left, right = st.columns(2, gap="large")

with left:
    st.subheader("Your text or files")
    st.text_area("Paste text", key="text", height=220,
                 placeholder="Paste your notes or article here…", label_visibility="collapsed")

    files = st.file_uploader(
        "Or upload files (PDF, Word, PowerPoint, Excel, CSV, TXT, images and more)",
        type=FILE_TYPES, accept_multiple_files=True, key=f"up_{st.session_state.up_key}")

    extracted = []
    for f in files or []:
        try:
            with st.spinner(f"Reading {f.name}…"):
                t = extract_text(f.name, f.getvalue())
            if t.strip():
                extracted.append(t)
                st.caption(f"✅ {f.name}: {len(words_of(t)):,} words found")
            else:
                st.warning(f"{f.name}: no readable text found. Scanned PDFs are not supported yet; "
                           f"upload them as images instead.")
        except Exception as e:
            st.warning(f"{f.name}: could not be read ({e})")

    full_text = "\n\n".join([st.session_state.text] + extracted).strip()
    n_words = len(words_of(full_text))
    st.caption(f"Total: {n_words:,} words · {max(1 if n_words else 0, round(n_words / 220))} min read")

    b1, b2, _ = st.columns([1, 1, 2])
    b1.button("Try sample", on_click=load_sample, use_container_width=True)
    b2.button("Clear", on_click=clear_all, use_container_width=True)

    level = st.select_slider("Summary length", options=list(LEVELS), value="Medium")
    fmt = st.radio("Format", ["Paragraph", "Bullet points"], horizontal=True)
    engines = ["Quick (offline)"] + (["AI (smarter)"] if ai_available() else [])
    engine = st.radio("Engine", engines, horizontal=True)

    go = st.button("Summarize", type="primary", use_container_width=True)

if go:
    if n_words < MIN_WORDS:
        st.session_state.result = {"error": f"Add more text or upload a file. At least {MIN_WORDS} words are needed."}
    else:
        try:
            with st.spinner("Reading your text…"):
                sents = (ai_summarize(full_text, level, fmt) if engine.startswith("AI")
                         else local_summarize(full_text, level))
            if not sents:
                raise ValueError("Could not find enough sentences. Try longer text with full sentences.")
            st.session_state.result = {"sents": sents, "fmt": fmt, "orig": full_text}
        except Exception as e:
            st.session_state.result = {"error": f"Something went wrong: {e}"}

with right:
    st.subheader("Summary")
    r = st.session_state.result
    if not r:
        st.info("Your summary will appear here. Paste text or upload a file, then select **Summarize**.")
    elif "error" in r:
        st.error(r["error"])
    else:
        sents = r["sents"]
        if r["fmt"] == "Bullet points":
            body = "<ul>" + "".join(f"<li>{htmllib.escape(s)}</li>" for s in sents) + "</ul>"
            plain = "\n".join("• " + s for s in sents)
        else:
            body = htmllib.escape(" ".join(sents))
            plain = " ".join(sents)
        st.markdown(f'<div class="summary-box">{body}</div>', unsafe_allow_html=True)

        in_w, out_w = len(words_of(r["orig"])), len(words_of(" ".join(sents)))
        m1, m2, m3 = st.columns(3)
        m1.metric("Words in summary", f"{out_w:,}")
        m2.metric("Shorter by", f"{max(0, round((1 - out_w / max(in_w, 1)) * 100))}%")
        m3.metric("Key sentences", len(sents))

        st.markdown("**Key terms**")
        st.markdown("".join(f'<span class="chip">{htmllib.escape(k)}</span>' for k in top_keywords(r["orig"])),
                    unsafe_allow_html=True)
        st.download_button("Download .txt", plain, file_name="study-sparks-summary.txt")
        with st.expander("Copy summary"):
            st.code(plain, language=None)

st.markdown(f"<hr><center style='color:#4B466F'>© Study Sparks AI · Created by {OWNER_NAME}</center>",
            unsafe_allow_html=True)
